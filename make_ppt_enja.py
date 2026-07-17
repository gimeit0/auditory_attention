"""
基于 7_ja_v2.pptx 的 13 页结构, 先做英文版(准确/自然), 再由英文翻成日语版。
输出新文件: presentation_en.pptx (Arial) 和 presentation_ja.pptx (Meiryo)。
含 Fig 2b 图页与音频页(嵌入 wav)。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

COL = {"dark": RGBColor(0x00, 0x00, 0x00), "grey": RGBColor(0x55, 0x55, 0x55),
       "red": RGBColor(0xc0, 0x39, 0x2b), None: None}
AUDIO = "export_audio/t0000_same"


def _set(run, size, font, bold=False, color=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def build(slides, font, outfile, audio_labels):
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    BLANK = prs.slide_layouts[6]
    for spec in slides:
        kind = spec[0]; s = prs.slides.add_slide(BLANK)
        if kind == "title":
            _, title, subs = spec
            tb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.4))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = title; _set(r, 32, font, True, COL["dark"])
            tb2 = s.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.7), Inches(1.8))
            for i, line in enumerate(subs):
                p = tb2.text_frame.paragraphs[0] if i == 0 else tb2.text_frame.add_paragraph()
                r = p.add_run(); r.text = line; _set(r, 15, font, color=COL["grey"])
        elif kind == "bullets":
            _, title, bullets = spec
            tb = s.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(1.0))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = title; _set(r, 25, font, True, COL["dark"])
            body = s.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(11.7), Inches(5.6))
            tf = body.text_frame; tf.word_wrap = True
            for i, item in enumerate(bullets):
                text, level = item[0], item[1]; ckey = item[2] if len(item) > 2 else None
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.level = level; p.space_after = Pt(5)
                mark = "• " if level == 0 else "– "
                r = p.add_run(); r.text = mark + text
                _set(r, 18 - level*2, font, bold=(level == 0 and ckey is not None), color=COL[ckey])
        elif kind == "image":
            _, title, img, caption = spec
            tb = s.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.9))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = title; _set(r, 23, font, True, COL["dark"])
            s.shapes.add_picture(img, Inches(0.6), Inches(1.5), height=Inches(5.4))
            cap = s.shapes.add_textbox(Inches(7.9), Inches(1.7), Inches(5.1), Inches(5.2))
            tf = cap.text_frame; tf.word_wrap = True
            for i, (text, ckey) in enumerate(caption):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(8)
                r = p.add_run(); r.text = text; _set(r, 14, font, color=COL[ckey])
        elif kind == "audio":
            _, title, subtitle = spec
            tb = s.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.8))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = title; _set(r, 23, font, True, COL["dark"])
            sb = s.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(12), Inches(0.7))
            sb.text_frame.word_wrap = True
            r = sb.text_frame.paragraphs[0].add_run(); r.text = subtitle; _set(r, 14, font, color=COL["grey"])
            grid = [("mix_-9dB.wav", 1.0, 2.5), ("mix_-6dB.wav", 3.0, 2.5), ("mix_-3dB.wav", 5.0, 2.5),
                    ("mix_+0dB.wav", 7.0, 2.5), ("mix_+3dB.wav", 9.0, 2.5), ("mix_inf.wav", 11.0, 2.5),
                    ("cue.wav", 3.0, 5.0), ("target_clean.wav", 8.0, 5.0)]
            for (fn, x, y), lab in zip(grid, audio_labels):
                s.shapes.add_movie(f"{AUDIO}/{fn}", Inches(x), Inches(y), Inches(1.0), Inches(1.0),
                                   mime_type="audio/wav")
                lb = s.shapes.add_textbox(Inches(x - 0.5), Inches(y + 1.02), Inches(2.0), Inches(0.4))
                lb.text_frame.word_wrap = True
                p = lb.text_frame.paragraphs[0]; p.alignment = 2
                r = p.add_run(); r.text = lab; _set(r, 12, font, color=COL["dark"])
    prs.save(outfile)
    print("已生成:", outfile, "(%d 页)" % len(prs.slides._sldIdLst))


# ============================= ENGLISH =============================
EN = [
 ("title", "Reproducing the Behavioral Trends of a Selective Auditory Attention Model",
  ["Reproduction target: Griffith, Hess & McDermott (2026), Nature Human Behaviour",
   "\"Optimized feature gains explain and predict human selective listening\"",
   "Presenter: [Your name]      Date: 2026-07"]),
 ("bullets", "Purpose of This Presentation", [
   ("This talk reports whether the authors' pretrained model (Griffith et al., 2026) can reproduce", 0),
   ("the behavioral trend of Fig 2b on a Common Voice evaluation set that I built myself", 1),
   ("Two questions this talk answers:", 0, "dark"),
   ("① Can the model actually recognize the target talker's middle word within a mixture?", 1),
   ("② Do the accuracy trends (vs SNR, and same- vs different-sex masker) match the paper's direction?", 1),
   ("Scope: no retraining — the authors' parameters are used as-is; we assess the trend direction, not exact values", 0, "red"),
   ("Why succeeding matters:", 0, "dark"),
   ("Shows the authors' model generalizes to an unseen corpus (Common Voice) — evidence of robustness", 1),
   ("Confirms my evaluation pipeline works — a foundation for future extensions", 1),
   ("Independently reproduces the paper's key behavioral findings — adds replication value", 1)]),
 ("bullets", "1. Research Background and Reproduction Goal", [
   ("Paper: a neural network with a 'feature-gain' attention mechanism that models human 'cocktail-party' listening", 0),
   ("Input a cue (a clip of the target talker) and a target+masker mixture → report the target's middle word in the mixture", 1),
   ("Guided by the cue, it amplifies features matching the target's timbre and suppresses the rest, separating the target", 1),
   ("Fig 2 shows, via behavioral curves, what governs success/failure (higher SNR → higher accuracy; different-sex maskers separate more easily)", 1),
   ("My reproduction goal:", 0, "dark"),
   ("① Validate the model: confirm the authors' pretrained model recognizes the middle word on my evaluation audio", 1),
   ("② Reproduce the Fig 2b trend: accuracy vs SNR under same- / different-sex maskers", 1),
   ("Criterion: success = the trend direction matches the paper; exact numbers are not expected (different corpus / talkers)", 1, "red")]),
 ("bullets", "1 (suppl.) How My Method Compares to the Authors' — What I Did, and How", [
   ("Same as the authors (used directly):", 0, "dark"),
   ("The model and pretrained parameters / task definition (cue + mixture → middle word) / evaluation framework (SNR, same/different sex) / scoring rule", 1),
   ("Different from the authors = work I did myself:", 0, "dark"),
   ("The authors' input data (training / evaluation stimuli) is not public → I had to build the evaluation set from scratch", 1, "red"),
   ("① Selected talkers and recordings from Common Voice v9 (sex-balanced, speaker-disjoint)", 1),
   ("② Estimated each word's time boundaries via forced alignment (information Common Voice lacks)", 1),
   ("③ Chose target words (anchors) using the 800-word vocabulary and timing-margin conditions", 1),
   ("④ Built cue / target / masker triplets, mixed at 6 SNRs (within-item), RMS-normalized, made diotic", 1),
   ("Key point: I did not simply reuse the authors' data — I validated their model on an evaluation set I built myself", 0, "red")]),
 ("bullets", "2. Data Preparation", [
   ("Corpus: Common Voice 9.0 English (the same version as the paper)", 0),
   ("Selection: talkers with a consistent sex label and ≥2 recordings → usable pool of 18,067 talkers / 971,848 recordings", 1),
   ("Target vocabulary: the authors' public 800 words (byte-identical to the official file); the model is an 800-class classifier", 0),
   ("Final evaluation set: 600 valid target samples", 0, "dark"),
   ("Strictly sex-balanced: 300 male / 300 female", 1),
   ("From 149 talkers and 347 recordings, covering 369 distinct target words", 1),
   ("600 samples keep each of the 6 SNR estimates stable (binomial SE ≈ ±2%)", 1)]),
 ("bullets", "3. Method Pipeline", [
   ("Raw audio → forced alignment → text normalization → anchor selection → cut & mix → model input → scoring", 0, "dark"),
   ("1. Forced alignment (MMS_FA, 16 kHz): obtain each word's time boundaries (CV has only sentence-level transcripts, no word timing)", 0),
   ("2. Text normalization: follows the paper's convert_transcript (lowercase, strip punctuation, spell out digits)", 0),
   ("3. Anchor (middle-word) selection: centered on the target word. Conditions = in vocabulary, word length < 2 s, ≥1.25 s margin on each side", 0),
   ("4. Cut & mix: each clip 2.5 s @ 44.1 kHz (model uses the central 2 s internally); within-item remixing at 6 SNRs", 0),
   ("SNR: -9/-6/-3/0/+3/+inf dB; masker scaled to the target's energy; +inf = no masker; RMS-normalized to 0.02; diotic (identical in both ears)", 1),
   ("5. Inference directly with the authors' pretrained checkpoint (word_task_v10); no retraining or fine-tuning", 0)]),
 ("bullets", "3 (cont.) Three Key Engineering Decisions — \"Why I Did It This Way\"", [
   ("① Keep apostrophes during text normalization", 0, "dark"),
   ("The paper's public code strips apostrophes, but the 800-word vocabulary contains 16 contractions (can't/don't, ...); stripping them causes lookup failures", 1),
   ("② Cut centered on the target word, not on the word at the recording's midpoint", 0, "dark"),
   ("Confirmed in the paper's source code: first choose the target word, then cut centered on it (not cut first, then check the midpoint word)", 1),
   ("③ Anchor-first pairing (align first, then pair)", 0, "dark"),
   ("Align only target candidates and keep recordings that have an anchor: 600 samples from 627 recordings (96%), far above the old method's 22%", 1)]),
 ("bullets", "4. Result: Model Validity", [
   ("Middle-word recognition accuracy under the clean condition (target only, no masker):", 0),
   ("88.6%   (chance level is only 0.1% = 1/800)", 0, "dark"),
   ("The +inf condition of the full 600 samples also gives 88.0%, corroborating this", 1),
   ("→ The model genuinely recognizes the target's middle word, far above chance", 0, "red"),
   ("This is the prerequisite for all subsequent trend experiments", 1, "grey")]),
 ("image", "4. Result: Fig 2b Reproduction — Same- vs Different-sex Masker", "fig2b_reproduction_en.png", [
   ("Two lines (same / different sex): middle-word accuracy vs SNR", "dark"),
   ("Reproduced trends (consistent with the paper):", "red"),
   ("① Accuracy increases monotonically with SNR (weaker masker → higher accuracy)", "grey"),
   ("② Different-sex masker is higher overall than same-sex (by 9–18 points)", "grey"),
   ("③ At +inf (no masker) the two lines converge to ~88%: sex is irrelevant with no masker", "grey"),
   ("Different > same is stable: error bars do not overlap at all 5 masked SNRs (3.4–6.8σ)", "dark"),
   ("Error bars = ±1 SEM", "grey")]),
 ("audio", "4 (suppl.) Audio Samples — SNR Ladder (click to play)",
  "The same stimulus (target word \"around\") mixed at 6 SNRs. From -9 dB to no masker, the masker talker gradually fades and the target becomes clear."),
 ("bullets", "4 (suppl.) Answers to the Two Questions of This Talk", [
   ("Q①: Can the model recognize the target talker's middle word?", 0, "dark"),
   ("→ Yes. 88.6% under clean (chance 0.1%); in mixtures it works consistently, from 26% at -9 dB to 88% at +inf", 1),
   ("Q②: Do the accuracy trends (vs SNR, masker sex) match the paper?", 0, "dark"),
   ("→ Yes. Accuracy rises with SNR (monotonic); different-sex > same-sex (non-overlapping error bars at every SNR); converges with no masker", 1),
   ("Both answers are affirmative → the behavioral trends match the paper's direction; the reproduction succeeded", 0, "red")]),
 ("bullets", "5. Key Judgments and Honest Caveats", [
   ("Strictly followed the paper: 800-word vocabulary, alignment approach, target-centered cutting, 6 SNRs, within-item, scoring rule, authors' parameters", 0),
   ("Adjustments made for practical reasons (the scientific setup is unchanged):", 0, "dark"),
   ("Keeping apostrophes (for contraction lookup) / anchor-first pairing (efficiency) / running on an M1 CPU", 1),
   ("Current scope and limits:", 0, "dark"),
   ("Scope is single-talker speech maskers, diotic (non-spatial). Other masker types, harmonicity, and spatial experiments are not done", 1),
   ("600 samples; one scan on CPU takes ~6.5 h. Can be scaled up on the hakusan GPU later", 1),
   ("Numbers differing from the paper is expected: corpus / talkers / words differ; we judge only the trend direction", 1, "red")]),
 ("bullets", "Summary", [
   ("Purpose: using the authors' pretrained model and my own CV evaluation set, test whether Fig 2b's behavioral trend can be reproduced", 0),
   ("Method: forced alignment → target-centered cutting → mixing at 6 SNRs → inference with the authors' checkpoint (600 samples, sex-balanced)", 0),
   ("Results:", 0, "dark"),
   ("① Model validity confirmed: clean accuracy 88.6% ≫ chance 0.1%", 1),
   ("② Fig 2b trend reproduced: accuracy↑ with SNR, different-sex > same-sex, convergence with no masker", 1),
   ("Conclusion: the trend direction matches the paper → reproduction succeeded (a trend reproduction, not a re-verification of the paper's truth)", 0, "red"),
   ("Next steps:", 0, "dark"),
   ("① Next time, train the model from scratch instead of using the pretrained checkpoint, and test whether the same effect (trend) is reproduced", 1, "red"),
   ("② Scale up on the hakusan GPU (more samples; add panels such as Fig 2c / 2d / 2e)", 1),
   ("③ Make maskers symmetric to bring the absolute confusion rate closer to the paper", 1),
   ("④ (Extension) Move toward validating own extensions such as the developmental parameter α", 1)]),
]
EN_AUDIO = ["-9 dB", "-6 dB", "-3 dB", "0 dB", "+3 dB", "+inf (no masker)", "cue", "target only (ref.)"]

# ======================= JAPANESE (from EN) =======================
JA = [
 ("title", "選択的聴覚注意モデルの行動的傾向の再現",
  ["再現対象論文: Griffith, Hess & McDermott (2026), Nature Human Behaviour",
   "\"Optimized feature gains explain and predict human selective listening\"",
   "発表者: [氏名]      日付: 2026-07"]),
 ("bullets", "本発表の目的", [
   ("本発表では、著者(Griffith et al., 2026)が公開した事前学習モデルを用い、", 0),
   ("自作したCommon Voice評価セット上で、Fig 2bの行動的傾向を再現できるかを報告する", 1),
   ("本発表で答える2つの問い:", 0, "dark"),
   ("① モデルは混合音声中の目標話者の中央単語を実際に認識できるか", 1),
   ("② 正答率の傾向(SN比、および同性/異性妨害)は、論文と方向が一致するか", 1),
   ("位置づけ: 再学習は行わず、著者のパラメータをそのまま使用。数値ではなく傾向の方向を評価する", 0, "red"),
   ("再現が成功する意義:", 0, "dark"),
   ("著者のモデルが未知のコーパス(Common Voice)にも汎化することを示せる —— 頑健性の証拠", 1),
   ("自作の評価パイプラインが正しく機能することを確認 —— 今後の拡張の基盤となる", 1),
   ("論文の主要な行動的知見を独立に再現し、追試としての価値を加える", 1)]),
 ("bullets", "一、研究背景と再現目標", [
   ("論文: 「特徴ゲイン」型の注意機構をもつニューラルネットワークで、人間の「カクテルパーティー」聴取を模擬する", 0),
   ("手がかり音(cue)と「目標+妨害」の混合音声を入力し、混合音声中の目標話者の中央単語を答える", 1),
   ("cueに導かれて目標の音色に合う特徴を増強し、それ以外を抑制することで目標を分離する", 1),
   ("Fig 2は認識の成否を規定する要因を行動曲線で示す(SN比が高いほど正答率↑、異性妨害ほど分離しやすい)", 1),
   ("本再現の目標:", 0, "dark"),
   ("① モデルの検証: 著者の事前学習モデルが、自作の評価音声上で中央単語を認識できることを確認する", 1),
   ("② Fig 2bの傾向の再現: 同性/異性妨害下での正答率のSN比依存性を再現する", 1),
   ("判定基準: 傾向の方向が論文と一致すれば成功。数値の一致は期待しない(コーパス・話者が異なるため)", 1, "red")]),
 ("bullets", "一(補)、著者の手法との比較 —— 私が何を、どう行ったか", [
   ("著者と同じ点(そのまま使用):", 0, "dark"),
   ("モデルと事前学習パラメータ / タスク定義(cue+混合→中央単語) / 評価の枠組み(SN比、同性/異性) / 採点規則", 1),
   ("著者と異なる点 = 私自身が行った作業:", 0, "dark"),
   ("著者の入力データ(学習・評価用の刺激)は非公開 → 評価セットを一から構築する必要があった", 1, "red"),
   ("① Common Voice v9 から話者と録音を選定(性別均衡・話者分離)", 1),
   ("② 強制アライメントで各単語の時間境界を推定(Common Voiceには無い情報)", 1),
   ("③ 800語彙と時間余白の条件で目標単語(アンカー)を選定", 1),
   ("④ cue/目標/妨害の三つ組を構成し、6段階SN比で混合(within-item)、RMS正規化、diotic化", 1),
   ("要点: 著者のデータを流用したのではなく、著者のモデルを自作の評価セットで検証した", 0, "red")]),
 ("bullets", "二、データ準備", [
   ("コーパス: Common Voice 9.0 英語版(論文と同一バージョン)", 0),
   ("選定: 性別ラベルが一貫し、かつ2録音以上をもつ話者 → 利用可能な 18,067名 / 971,848録音", 1),
   ("目標語彙: 著者公開の800語(公式ファイルとバイト単位で一致)。モデルは800クラスの分類器", 0),
   ("最終評価セット: 有効な目標サンプル 600件", 0, "dark"),
   ("厳密に性別均衡: 男性300 / 女性300", 1),
   ("149名の話者・347録音に由来し、369種の異なる目標単語を含む", 1),
   ("600サンプルにより、6段階のSN比それぞれで推定が安定(二項標準誤差 約±2%)", 1)]),
 ("bullets", "三、手法パイプライン", [
   ("原音声 → 強制アライメント → テキスト正規化 → アンカー選定 → 切り出し・混合 → モデル入力 → 採点", 0, "dark"),
   ("1. 強制アライメント(MMS_FA, 16 kHz): 各単語の時間境界を取得(CVは文単位の書き起こしのみで、単語の時刻をもたない)", 0),
   ("2. テキスト正規化: 論文のconvert_transcriptに準拠(小文字化・記号除去・数字の単語化)", 0),
   ("3. アンカー(中央単語)選定: 目標単語を中心とする。条件=語彙内、単語長<2s、前後に各1.25s以上の余白", 0),
   ("4. 切り出し・混合: 各2.5s@44.1kHz(モデルは内部で中央2sを使用)。within-itemで6段階のSN比に再混合", 0),
   ("SN比: -9/-6/-3/0/+3/+inf dB。目標のエネルギーに合わせ妨害を調整。+inf=妨害なし。RMS 0.02に正規化。diotic(両耳同一)", 1),
   ("5. 著者の事前学習チェックポイント(word_task_v10)でそのまま推論。再学習・微調整は行わない", 0)]),
 ("bullets", "三(続)、3つの重要な工学的判断 —— 「なぜそうしたか」", [
   ("① テキスト正規化でアポストロフィを保持する", 0, "dark"),
   ("論文の公開コードはアポストロフィを削除するが、800語彙には16個の短縮形(can't/don't 等)が含まれ、削除すると照合に失敗する", 1),
   ("② 録音の中点の単語ではなく、目標単語を中心に切り出す", 0, "dark"),
   ("論文のソースコードで確認: まず目標単語を選び、それを中心に切り出す(先に切り出して中点の単語を見るのではない)", 1),
   ("③ アンカー優先のペアリング(先にアライメント、後にペアリング)", 0, "dark"),
   ("目標候補のみを整列し、アンカーをもつ録音を採用: 627録音から600サンプル(96%)。旧手法の22%を大きく上回る", 1)]),
 ("bullets", "四、結果: モデルの有効性", [
   ("clean条件(目標のみ・妨害なし)における中央単語認識の正答率:", 0),
   ("88.6%   (偶然水準はわずか0.1% = 1/800)", 0, "dark"),
   ("600サンプルの+inf条件でも88.0%となり、これを裏付ける", 1),
   ("→ モデルは目標の中央単語を確かに認識でき、偶然水準を大きく上回る", 0, "red"),
   ("これが以降の傾向実験すべての前提となる", 1, "grey")]),
 ("image", "四、結果: Fig 2b 再現 —— 同性 vs 異性妨害", "fig2b_reproduction_en.png", [
   ("2本の線(同性/異性): 中央単語認識の正答率 vs SN比", "dark"),
   ("再現された傾向(論文と一致):", "red"),
   ("① SN比とともに正答率が単調に増加(妨害が弱いほど正答率↑)", "grey"),
   ("② 異性妨害が全体的に同性より高い(9~18ポイント)", "grey"),
   ("③ +inf(妨害なし)で2本の線が約88%に収束: 妨害がなければ性別は無関係", "grey"),
   ("異性>同性は安定: 妨害ありの5段階すべてでエラーバーが非重複(3.4~6.8σ)", "dark"),
   ("エラーバー = ±1 SEM", "grey")]),
 ("audio", "四(補)、音声サンプル —— SN比ラダー(クリックで再生)",
  "同一刺激(目標単語「around」)を6段階のSN比で混合。-9dBから妨害なしへ進むと、妨害話者が徐々に消え目標が明瞭になる。"),
 ("bullets", "四(補)、本発表の2つの問いへの回答", [
   ("問い①: モデルは目標話者の中央単語を認識できるか?", 0, "dark"),
   ("→ できる。clean条件で88.6%(偶然0.1%)。混合音声下でも -9dBの26% 〜 +infの88% まで一貫して機能する", 1),
   ("問い②: 正答率の傾向(SN比・妨害の性別)は論文と一致するか?", 0, "dark"),
   ("→ 一致する。SN比↑で正答率↑(単調)、異性>同性(全SN比でエラーバー非重複)、妨害なしで収束", 1),
   ("いずれも肯定 → 行動的傾向は論文と方向が一致し、再現に成功した", 0, "red")]),
 ("bullets", "五、重要な判断と誠実な補足", [
   ("論文を厳密に踏襲した点: 800語彙、アライメント方針、目標単語中心の切り出し、6段階SN比、within-item、採点規則、著者パラメータ", 0),
   ("実情に応じて調整した点(科学的な設定は変更しない):", 0, "dark"),
   ("アポストロフィの保持(短縮形照合のため)/ アンカー優先ペアリング(効率化)/ M1 CPU上での実行", 1),
   ("現在の範囲と限界:", 0, "dark"),
   ("範囲は単一話者の音声妨害・diotic(非空間)。他の妨害種・harmonicity・空間実験は未実施", 1),
   ("サンプル600件、CPUでの1回の走査は約6.5時間。今後hakusan GPUで大規模化できる", 1),
   ("数値が論文と異なるのは想定内: コーパス・話者・単語が異なるため。判定はあくまで傾向の方向", 1, "red")]),
 ("bullets", "まとめ", [
   ("目的: 著者の事前学習モデルと自作CV評価セットを用い、Fig 2bの行動的傾向を再現できるかを検証した", 0),
   ("方法: 強制アライメント → 目標単語中心の切り出し → 6段階SN比で混合 → 著者チェックポイントで推論(600サンプル、性別均衡)", 0),
   ("結果:", 0, "dark"),
   ("① モデルの有効性を確認: clean正答率 88.6% ≫ 偶然 0.1%", 1),
   ("② Fig 2bの傾向を再現: SN比↑で正答↑、異性>同性、妨害なしで収束", 1),
   ("結論: 傾向の方向が論文と一致 → 再現に成功(傾向の再現であり、論文の真偽検証ではない)", 0, "red"),
   ("次にやること:", 0, "dark"),
   ("① 次回は学習済みチェックポイントを使わず、モデルを自分で一から学習し、同じ効果(傾向)が再現されるかを検証する", 1, "red"),
   ("② hakusan GPUで大規模化(サンプル増、Fig 2c/2d/2e など他パネルの追加)", 1),
   ("③ 妨害を対称構成にし、混同の絶対値を論文水準に近づける", 1),
   ("④ (発展) 発達パラメータαなど、独自拡張の検証へ進む", 1)]),
]
JA_AUDIO = ["-9 dB", "-6 dB", "-3 dB", "0 dB", "+3 dB", "+inf(妨害なし)", "cue", "目標のみ(参照)"]

build(EN, "Arial", "presentation_en.pptx", EN_AUDIO)
build(JA, "Meiryo", "presentation_ja.pptx", JA_AUDIO)
