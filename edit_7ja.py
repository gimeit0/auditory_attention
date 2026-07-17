"""
就地编辑 7_ja.pptx(用户手改版, Hiragino Sans), 落实导师4条建议, 输出 7_ja_v2.pptx。
不覆盖原文件(原文件正被打开)。新内容用 Hiragino Sans 匹配该文件风格。
  A. 背景页后 新增「著者の手法との異同」
  B. 目的页 追加「成功する意義」
  C. まとめ页 追加「次にやること」
  D. 音声サンプル页后 新增「2つの問いへの回答」
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

FONT = "Hiragino Sans"
BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xc0, 0x39, 0x2b)
GREY = RGBColor(0x55, 0x55, 0x55)

prs = Presentation("7_ja.pptx")


def style(run, size, color=None, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def body_of(slide):
    """取正文文本框(top>1in, 排除标题)。"""
    cands = [sh for sh in slide.shapes
             if sh.has_text_frame and sh.top is not None and sh.top / 914400 > 1.0]
    return max(cands, key=lambda sh: sh.height)


def append_bullets(tf, bullets):
    for item in bullets:
        text, level = item[0], item[1]
        color = item[2] if len(item) > 2 else None
        p = tf.add_paragraph()
        p.level = level
        p.space_after = Pt(6)
        mark = "• " if level == 0 else "– "
        r = p.add_run(); r.text = mark + text
        style(r, 19 - level * 2, color, bold=(level == 0 and color is not None))


def new_bullets_slide(title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(1.0))
    r = tb.text_frame.paragraphs[0].add_run(); r.text = title
    style(r, 26, BLACK, bold=True)
    body = s.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(11.6), Inches(5.5))
    tf = body.text_frame; tf.word_wrap = True
    for i, item in enumerate(bullets):
        text, level = item[0], item[1]
        color = item[2] if len(item) > 2 else None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level; p.space_after = Pt(6)
        mark = "• " if level == 0 else "– "
        r = p.add_run(); r.text = mark + text
        style(r, 19 - level * 2, color, bold=(level == 0 and color is not None))
    return s


def move_slide(to_index):
    """把最后添加的幻灯片移动到 to_index(0基)。"""
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert(to_index, el)


# ---- B. 目的页 追加「成功する意義」 (slide index 1) ----
append_bullets(body_of(prs.slides[1]).text_frame, [
    ("本再現が成功する意義:", 0, BLACK),
    ("著者モデルが未知のコーパス(Common Voice)にも汎化することを示せる —— モデルの頑健性・一般性の傍証", 1),
    ("自作の評価パイプラインが正しく機能することを確認 —— 今後の拡張(他条件・独自研究)の基盤になる", 1),
    ("論文の主要な行動的知見が独立に再現され、知見の信頼性(追試価値)を高める", 1),
])

# ---- C. まとめ页 追加「次にやること」 (最后一页) ----
matome = prs.slides[len(prs.slides._sldIdLst) - 1]
mb = body_of(matome)
mb.height = Inches(5.3)   # 扩高以容纳新增行
append_bullets(mb.text_frame, [
    ("次にやること:", 0, BLACK),
    ("① hakusan GPUで大規模化(サンプル増、Fig 2c/2d/2e など他パネルの追加)", 1),
    ("② 妨害も対称構成にし、混同(confusion)の絶対量を論文水準に近づける", 1),
    ("③ (発展) 発達パラメータαなど、独自拡張の検証へ進む", 1),
])

# ---- A. 新增「著者の手法との異同」, 放到 背景(index2) 之后 = index3 ----
new_bullets_slide("一(補)、著者の手法との異同 —— 私は何を・どう行ったか", [
    ("著者と同じ点(そのまま踏襲):", 0, BLACK),
    ("モデル本体と学習済みパラメータ / タスク定義(cue+混合音声→中央単語) / 評価の枠組み(SN比・同性/異性) / 採点基準", 1),
    ("著者と異なる点 = 私が独自に行った作業:", 0, BLACK),
    ("著者の入力データ(学習・評価用の刺激音声)は非公開 → 評価セットを一から自作する必要があった", 1, RED),
    ("① Common Voice v9 から話者・録音を選別(性別均衡・話者分離)", 1),
    ("② 強制アライメントで各単語の時間境界を推定(CVには無い情報)", 1),
    ("③ 800語彙・時間余白の条件で目標単語(アンカー)を選定", 1),
    ("④ cue/目標/妨害の三つ組を構成し、6段階SN比で混合(within-item)・RMS正規化・diotic化", 1),
    ("要点: 「著者のデータをそのまま流用した」のではなく、著者のモデルを私が独自に構築した評価セットで検証した", 0, RED),
])
move_slide(3)

# ---- D. 新增「2つの問いへの回答」, 放到 音声サンプル 之后 ----
# 此时: 0 title,1 目的,2 背景,3 異同,4 データ,5 手法,6 決定,7 結果1,8 結果2,9 音声,10 五,11 まとめ
new_bullets_slide("四(補)、本発表の2つの問いへの回答", [
    ("問い①: モデルは目標話者の中央単語を認識できるか?", 0, BLACK),
    ("→ 認識できる。clean条件で正答率88.6%(偶然0.1%)。混合音声下でも -9dBの26% 〜 +infの88% まで一貫して機能", 1),
    ("問い②: 正答率のSN比・妨害話者性別への依存傾向は論文と一致するか?", 0, BLACK),
    ("→ 一致する。SN比↑で正答率↑(単調増加)、異性妨害>同性(全SN比でエラーバー非重複)、妨害なしで収束", 1),
    ("いずれの問いも肯定 → 行動的傾向は論文と方向が一致し、再現に成功", 0, RED),
])
move_slide(10)  # 音声(9) の直後

prs.save("7_ja_v2.pptx")
print("已生成: 7_ja_v2.pptx")
for i, s in enumerate(prs.slides):
    t = s.shapes[0].text_frame.text if s.shapes and s.shapes[0].has_text_frame else ""
    print(" %2d. %s" % (i + 1, t[:40]))
